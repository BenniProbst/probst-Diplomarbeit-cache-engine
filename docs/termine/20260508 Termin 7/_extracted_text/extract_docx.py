"""Extrahiert alle docx-Dateien zu Text fuer Quellenanalyse."""
import os
import sys
from pathlib import Path
from docx import Document

ROOT = Path(r"C:\Users\benja\OneDrive\Desktop\Diplomarbeit - Datenbanken")
OUT_DIR = ROOT / "20260508 Termin 7" / "_extracted_text"
OUT_DIR.mkdir(parents=True, exist_ok=True)

def extract_docx(path: Path) -> str:
    """Extrahiert Volltext aus docx inklusive Tabellen."""
    doc = Document(path)
    parts = []
    # Paragraphen
    for p in doc.paragraphs:
        text = p.text.strip()
        if text:
            parts.append(text)
    # Tabellen
    for ti, table in enumerate(doc.tables):
        parts.append(f"\n[TABELLE {ti+1}]")
        for ri, row in enumerate(table.rows):
            cells = [c.text.strip().replace("\n", " | ") for c in row.cells]
            parts.append(" || ".join(cells))
        parts.append("[/TABELLE]\n")
    return "\n".join(parts)

def safe_name(name: str) -> str:
    return "".join(c if c.isalnum() or c in "._- " else "_" for c in name)

count = 0
for docx_path in sorted(ROOT.rglob("*.docx")):
    # Skip Termin 7 selbst
    if "Termin 7" in str(docx_path):
        continue
    rel = docx_path.relative_to(ROOT)
    out_name = safe_name(str(rel).replace(os.sep, "__")) + ".txt"
    out_file = OUT_DIR / out_name
    try:
        text = extract_docx(docx_path)
        out_file.write_text(text, encoding="utf-8")
        size = len(text)
        count += 1
        print(f"OK [{count}]: {rel}  -> {out_name} ({size} chars)")
    except Exception as e:
        print(f"FEHLER: {rel}: {e}", file=sys.stderr)

print(f"\nGesamt: {count} Dateien extrahiert nach {OUT_DIR}")
